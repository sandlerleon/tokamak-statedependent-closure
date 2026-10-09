# -*- coding: utf-8 -*-
"""Cover letter for IEEE Transactions on Plasma Science (regular paper).   python build_cover_letter.py -> out/Cover_Letter_IEEE_TPS.docx"""
import json
import os
import sys

import numpy as np
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docx_helpers as H  # noqa: E402

ROOT = os.path.join(HERE, "..")
R = json.load(open(os.path.join(ROOT, "results.json"), encoding="utf-8"))
ZEN = json.load(open(os.path.join("C:" + os.sep, "YouTube", "_tok_zenodo_state.json")))
RELEASE = os.environ.get("RELEASE_TAG", "v1.2.0")
SW = (ZEN.get("software_" + RELEASE[1:]) or ZEN["software"])["doi"]
PP = (ZEN.get("publication_v4") or ZEN["publication"])["doi"]
REPO = "https://github.com/sandlerleon/tokamak-statedependent-closure"
TITLE = "Numerical Admissibility and Regularization of Shear-Suppression Closures for Reduced Tokamak Transport"
TB = {(r["sc"] if r["sc"] is not None else "base"): r for r in R["table1"]}
LC = R["local_closure"]
LR = R["linear_response"]
CST = {r["sc"]: {x["torque"]: x["Q"] for x in r["rows"]} for r in R["coupled_sc_torque"]}
EV = {r["name"]: r for r in R["edge_variants"]}
doc = H.new_document(size=11, line=1.15)


def para(text, bold=False, after=6):
    p = doc.add_paragraph()
    if bold:
        p.add_run(text).bold = True
    else:
        H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    H.add_rich(p, text)
    p.paragraph_format.space_after = Pt(3)


for line in ("Leon Sandler", "Independent researcher, Northbrook, Illinois, USA", "sandler.leon@gmail.com", "ORCID: https://orcid.org/0009-0007-4584-808X"):
    para(line, after=0)
para("")
para("9 October 2026")
para("The Editor-in-Chief\nIEEE Transactions on Plasma Science", after=10)
para("Submission of a regular paper: “%s”" % TITLE, bold=True, after=10)
para("Dear Editor,")
para("I submit the enclosed paper for consideration as a regular paper in the IEEE Transactions on Plasma Science. It addresses a modeling question that matters for fusion transport codes: when a turbulent-transport closure is suppressed by sheared E×B flow, "
     "does it define a numerically well-behaved steady problem, and how much does it change the fusion gain once it does? The study is computational and theoretical, uses a reduced one-dimensional model of a reactor-scale tokamak, and uses no experimental data.")
para("Main results", bold=True, after=3)
bullet("*Numerical admissibility and regularization.* The local closure, whose shearing rate contains the second derivative of the temperature, is numerically ill posed: its converged answer depends on how the edge is treated (%.0f%% apart at a threshold of 0.05 for two natural edge conditions), "
       "and a grid-scale instability appears at a threshold that grows as N^{%.2f} with the number of cells N, so that every finite threshold becomes unstable at fine resolution. A first-order edge treatment hides this and produces a spurious fold. "
       "An adaptive-field closure that smooths the shearing rate over a fixed length is stable and grid converged over the tested conditions, converges at second order, and is insensitive to the length (%.1f%% over a tenfold range)."
       % (100 * abs(EV["local, edge value held"]["Q@0.05/N400"] / EV["local, no suppression in the last cell"]["Q@0.05/N400"] - 1), LC["threshold_exponent"],
          100 * (max(EV["smoothed, l = %s m" % x]["Q@0.1/N400"] for x in ("0.02", "0.05", "0.10", "0.20")) / min(EV["smoothed, l = %s m" % x]["Q@0.1/N400"] for x in ("0.02", "0.05", "0.10", "0.20")) - 1)))
bullet("*Closed-form admissibility conditions* for a shear-dependent viscosity in the rotation equation (m ≤ 1 for 1/(1 + Λ^{m}); floors 1/9 and %.4f), with the predicted saturation, fold, and hysteresis window, "
       "reproduced by direct solutions to a relative error of %.0e." % (R["theory"]["thresholds"]["floor_exp"], R["scalar_flux_relation_summary"]["max_rel_err"]))
bullet("*A fitted gain relation and its uncertainty.* The smoothed closure raises the fusion gain by %.2f%% at a suppression threshold of 0.3 and %.1f%% at 0.1, following the fitted relation ΔQ/Q = %.3f%%/s_{c}² (an empirical fit over the tested range, not a universal scaling); a Sobol study over seven parameters gives the spread, and the baseline is consistent with the "
       "ITER89-P L-mode scaling (H_{89} = %.2f; with the edge temperature imposed this is a consistency check, not a validation)." % (100 * TB[0.3]["dQ"], 100 * TB[0.1]["dQ"], 100 * LR["C_mean"] / TB["base"]["Q"], TB["base"]["H89"]))
bullet("*Physical calibration and a screening-level plant balance* (Section VI) against confinement scalings, the Greenwald and Troyon limits, neutral-beam torque, and an assumed-efficiency electric power balance. At the 40 MW reference operating point the estimated net electric power is approximately %.0f MW and remains negative across all %d sampled plant-efficiency assumptions; a separate auxiliary-power scan finds positive net power for %.0f%% of the sampled assumptions at the most favorable low heating power. These are screening-level estimates, not predictions of a viable power plant; the closure adds about %.0f MW at a threshold of 0.1, and no cost is computed. At the estimated torque of a reactor-scale device (%.0f N m) the torque-driven gain is about %.1f%% (threshold 0.5)."
       % (R["plant"]["cases"][0]["Pnet"], R["plant"]["uncertainty"]["n"], 100 * R["plant"]["uncertainty"]["best_over_Paux_baseline"]["frac_positive"], R["plant"]["cases"][2]["dPnet"], R["calibration"]["nbi_torque"]["5.3"], 100 * (CST[0.5][36.0] / CST[0.5][0.0] - 1)))
para("Fit to the journal", bold=True, after=3)
para("The paper concerns transport modeling for fusion devices and the numerical reliability of the closures used in it; it states what has and has not been established and keeps the claims no stronger than the checks. “Well posed” is used in a numerical sense: no existence, uniqueness, or continuous-dependence theorem is claimed for the nonlinear heat problem, only the admissibility conditions of the rotation equation are proved. The Discussion also states, without claiming a design result, how the method bears on screening compact, cost-constrained tokamaks. "
     "A supplementary file contains the proofs, the numerical protocol, and extended tables.")
para("Code, data, and declarations", bold=True, after=3)
para("All code, tests, raw results, and figure scripts are public at %s (release %s) and archived at https://doi.org/%s; the manuscript is available as a preprint at https://doi.org/%s. Reproducing every number takes under an hour. "
     "In line with IEEE policy, the use of artificial intelligence is disclosed in the Acknowledgment: Claude Sonnet 5.5 (Anthropic; model identifier claude-sonnet-5-5), accessed through the Claude Code environment of the Claude desktop application, was used between 7 and 9 October 2026 "
     "for code, derivations, numerical experiments, figures, and drafting; I reviewed and edited all content, checked every reference against Crossref, and take full responsibility for the paper. I am the sole author, have no conflicts of interest, and received no funding. "
     "The manuscript has not been published in a journal and is not under consideration elsewhere. An earlier version of the preprint (https://doi.org/%s) reported a steady-state fold that I have since shown to be an artifact of a first-order edge treatment; it is superseded by the version cited above. Intermediate archives (versions 1.1.0 and 1.1.1, https://doi.org/10.5281/zenodo.23251966 and https://doi.org/10.5281/zenodo.23267849) are superseded by the cited version, which adds the plant power balance and qualifies the well-posedness and gain-law claims." % (REPO, RELEASE, SW, PP, ZEN["publication"]["doi"]))
para("Thank you for considering the paper.")
para("Yours sincerely,", after=18)
para("Leon Sandler")
out = os.path.join(HERE, "out", "Cover_Letter_IEEE_TPS.docx")
doc.save(out)
print("saved", out)
